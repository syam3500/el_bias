"""
generate_results_excel.py
--------------------------
Extracts the last summary-statistics table from every simulation .md file
across all four AET directories (0.45, 3, 6, 36 µg/mL), assembles a master
DataFrame with scenario metadata, and writes Results_by_scenario.xlsx with
one sheet per analytical platform.

Columns written
---------------
Platform | AET (µg/mL) | Surrogate | Conc. Distribution | RRF Distribution |
File | Method |
FN Rate | FN SD | FN 95% PI Lo | FN 95% PI Hi |
FP Rate | FP SD | FP 95% PI Lo | FP 95% PI Hi |
Recall | Precision | F1 | F5 | Applied UF |
% True Above AET | % Reported Above AET

95% Prediction Interval (PI)
------------------------------
PI = mean +/- 1.96 x SD  (captures 95% of individual Monte Carlo run outcomes;
R = 10 000 iterations, N = 80 compounds per run).
All rate/SD/PI values are expressed as percentages (x100 from raw [0,1] values).

Usage
-----
python generate_results_excel.py
"""

import os
import re
import glob
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE      = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.normpath(os.path.join(HERE, "..", ".."))
BASE_DIR  = os.path.join(REPO_ROOT, "results", "scenarios")
XLSX_PATH = os.path.join(REPO_ROOT, "results", "Results_by_scenario.xlsx")
AET_DIRS  = ["AET_0.45", "AET_3", "AET_6", "AET_36"]
Z95 = 1.96

_ROW_RE = re.compile(r"^(\s*(?:\.{3}\d+|\d+%))\s+(.*)")


def _parse_block(block):
    rows = {}
    hdr = []
    has_method = False
    for raw in block.split("\n"):
        s = raw.strip()
        if not s:
            continue
        m = _ROW_RE.match(s)
        if m:
            idx  = m.group(1).strip()
            rest = m.group(2).strip()
            parts = rest.split()
            if not hdr:
                continue
            if idx not in rows:
                rows[idx] = {}
            if has_method:
                n_num = len(hdr) - 1
                if len(parts) >= n_num:
                    rows[idx]["Method"] = " ".join(parts[:-n_num]) if n_num else " ".join(parts)
                    for col, val in zip(hdr[1:], parts[-n_num:]):
                        rows[idx][col] = val
            else:
                for col, val in zip(hdr, parts[-len(hdr):]):
                    rows[idx][col] = val
        else:
            cols = s.split()
            has_method = "Method" in cols
            hdr = cols
    return rows


def parse_summary_table(content):
    hits = list(re.finditer(r"^\s*Method\s+FN_Rate\s+FP_Rate", content, re.MULTILINE))
    if not hits:
        return None
    start = hits[-1].start()
    block = content[start:]
    end   = block.find("\n```")
    if end > 0:
        block = block[:end]
    rows = _parse_block(block)
    if not rows:
        return None
    df = pd.DataFrame.from_dict(rows, orient="index")
    cols = ["Method"] + [c for c in df.columns if c != "Method"]
    return df[[c for c in cols if c in df.columns]]


def get_metadata(md_path):
    fname  = os.path.basename(md_path)
    base   = fname.replace(".md", "")
    folder = os.path.basename(os.path.dirname(md_path))
    m   = re.match(r"AET[\s_]+([\d.]+)", folder, re.IGNORECASE)
    aet = float(m.group(1)) if m else None
    if   "LC_MS+_" in base: platform = "LC/MS POS"
    elif "LC_MS-_" in base: platform = "LC/MS NEG"
    elif "LC_MS_X"  in base: platform = "LC/MS (Sim.)"
    elif "GC_MS"    in base: platform = "GC/MS"
    else:                     platform = "Unknown"
    for pfx in ("LC_MS+_", "LC_MS-_", "LC_MS_", "GC_MS_"):
        if base.startswith(pfx):
            rest = base[len(pfx):]
            break
    else:
        rest = base
    parts = rest.split("_")
    is_lc_sim = "LC_MS_X" in base
    if is_lc_sim:
        surrogate = "N/A"
        conc_rrf  = parts[1] if len(parts) > 1 else ""
        suffix    = "_".join(parts[2:]) if len(parts) > 2 else ""
    else:
        surr_raw  = parts[0]
        conc_rrf  = parts[1] if len(parts) > 1 else ""
        suffix    = "_".join(parts[2:]) if len(parts) > 2 else ""
        surrogate = "N/A" if surr_raw == "X" else f"{surr_raw} \u00b5g/mL"
    conc_id = conc_rrf[0] if conc_rrf else ""
    rrf_id  = conc_rrf[1] if len(conc_rrf) > 1 else ""
    if conc_id == "1":
        conc_dist = {"0.60": "Uniform 40%", "0.70": "Uniform 30%"}.get(suffix, "Uniform 66.7%")
    elif conc_id == "2":
        conc_dist = {"GC": "Log-Normal (Fitted GC)", "LC": "Log-Normal (Fitted LC)"}.get(suffix, "Log-Normal Standard")
    elif conc_id == "3":
        conc_dist = "Custom Mixture"
    else:
        conc_dist = "Simulated (Normal)"
    rrf_dist = {"A": "Empirical (CLAP)", "B": "Simulated (Normal)"}.get(rrf_id, "Unknown")
    return {
        "Platform":           platform,
        "AET (\u00b5g/mL)":        aet,
        "Surrogate":          surrogate,
        "Conc. Distribution": conc_dist,
        "RRF Distribution":   rrf_dist,
        "File":               base + ".md",
    }


def numeric(val):
    try:    return float(val)
    except: return None


def pi_bound(mean_pct, sd_pct, direction):
    if mean_pct is None or sd_pct is None:
        return None
    return round(mean_pct + direction * Z95 * sd_pct, 5)


UF_MAP = {"No Correction": 1, "UF (2)": 2, "UF (4)": 4, "UF (5)": 5, "UF (10)": 10}

all_rows = []

for aet_dir in AET_DIRS:
    folder_path = os.path.join(BASE_DIR, aet_dir)
    md_files    = glob.glob(os.path.join(folder_path, "*.md"))
    for md_path in sorted(md_files):
        fname = os.path.basename(md_path)
        if "readme" in fname.lower() or "combined" in fname.lower():
            continue
        with open(md_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        df = parse_summary_table(content)
        if df is None or df.empty:
            print(f"  SKIP (no table): {aet_dir}/{fname}")
            continue
        meta = get_metadata(md_path)
        aet_val = meta["AET (µg/mL)"]
        print(f"  OK  [{aet_val} µg/mL] {fname}  -  {len(df)} methods")
        for _, row in df.iterrows():
            method = str(row.get("Method", "")).strip()
            fn_r  = numeric(row.get("FN_Rate"))
            fp_r  = numeric(row.get("FP_Rate"))
            fn_sd = numeric(row.get("FN_SD"))
            fp_sd = numeric(row.get("FP_SD"))
            rec   = numeric(row.get("recall"))
            prec  = numeric(row.get("precision"))
            f1    = numeric(row.get("F1"))
            f5    = numeric(row.get("f_beta"))
            n_true  = numeric(row.get("is_at_and_above_threshold"))
            n_not   = numeric(row.get("X.is_at_and_above_threshold"))
            n_rep   = numeric(row.get("reported_at_and_above_threshold"))
            n_not_r = numeric(row.get("X.reported_at_and_above_threshold"))
            pct_true = (n_true / (n_true + n_not)  * 100) if (n_true is not None and n_not   is not None and (n_true + n_not)   > 0) else None
            pct_rep  = (n_rep  / (n_rep  + n_not_r) * 100) if (n_rep  is not None and n_not_r is not None and (n_rep  + n_not_r) > 0) else None
            applied_uf = UF_MAP.get(method)
            if applied_uf is None:
                m_uf = re.search(r"UF\s*\((\d+(?:\.\d+)?)\)", method)
                if m_uf:
                    applied_uf = float(m_uf.group(1))
            if applied_uf is None:
                applied_uf = numeric(row.get("UF"))
            fn_rate_pct = round(fn_r  * 100, 5) if fn_r  is not None else None
            fn_sd_pct   = round(fn_sd * 100, 7) if fn_sd is not None else None
            fp_rate_pct = round(fp_r  * 100, 5) if fp_r  is not None else None
            fp_sd_pct   = round(fp_sd * 100, 7) if fp_sd is not None else None
            entry = {**meta, "Method": method}
            entry["FN Rate"]              = fn_rate_pct
            entry["FN SD"]                = fn_sd_pct
            entry["FN 95% PI Lo"]         = pi_bound(fn_rate_pct, fn_sd_pct, -1)
            entry["FN 95% PI Hi"]         = pi_bound(fn_rate_pct, fn_sd_pct, +1)
            entry["FP Rate"]              = fp_rate_pct
            entry["FP SD"]                = fp_sd_pct
            entry["FP 95% PI Lo"]         = pi_bound(fp_rate_pct, fp_sd_pct, -1)
            entry["FP 95% PI Hi"]         = pi_bound(fp_rate_pct, fp_sd_pct, +1)
            entry["Recall"]               = round(rec  * 100, 5) if rec  is not None else None
            entry["Precision"]            = round(prec * 100, 5) if prec is not None else None
            entry["F1"]                   = round(f1   * 100, 5) if f1   is not None else None
            entry["F5"]                   = round(f5   * 100, 5) if f5   is not None else None
            entry["Applied UF"]           = applied_uf
            entry["% True Above AET"]     = round(pct_true, 5) if pct_true is not None else None
            entry["% Reported Above AET"] = round(pct_rep,  5) if pct_rep  is not None else None
            all_rows.append(entry)

master = pd.DataFrame(all_rows)
print(f"\nTotal rows extracted: {len(master)}")
print(master["Platform"].value_counts().to_string())

COLS = [
    "Platform", "AET (\u00b5g/mL)", "Surrogate", "Conc. Distribution",
    "RRF Distribution", "File", "Method",
    "FN Rate", "FN SD", "FN 95% PI Lo", "FN 95% PI Hi",
    "FP Rate", "FP SD", "FP 95% PI Lo", "FP 95% PI Hi",
    "Recall", "Precision", "F1", "F5",
    "Applied UF", "% True Above AET", "% Reported Above AET",
]

SHEET_MAP = {
    "GC-MS":      master[master["Platform"] == "GC/MS"],
    "LC-MS POS":  master[master["Platform"] == "LC/MS POS"],
    "LC-MS NEG":  master[master["Platform"] == "LC/MS NEG"],
    "LC-MS Sim.": master[master["Platform"] == "LC/MS (Sim.)"],
    "All":        master,
}

HDR_FONT  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HDR_FILL  = PatternFill("solid", start_color="1F4E79")
PI_FILL   = PatternFill("solid", start_color="2E6FAE")
HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN      = Side(style="thin", color="BFBFBF")
CELL_BDR  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT_FILL  = PatternFill("solid", start_color="D6E4F0")
NUM_FMT    = "0.00000"
NUM_FMT_SD = "0.0000000"

PI_COLS = {"FN 95% PI Lo", "FN 95% PI Hi", "FP 95% PI Lo", "FP 95% PI Hi"}

COL_WIDTHS = {
    "Platform": 14, "AET (\u00b5g/mL)": 11, "Surrogate": 13,
    "Conc. Distribution": 22, "RRF Distribution": 18, "File": 28,
    "Method": 24,
    "FN Rate": 10, "FN SD": 12, "FN 95% PI Lo": 13, "FN 95% PI Hi": 13,
    "FP Rate": 10, "FP SD": 12, "FP 95% PI Lo": 13, "FP 95% PI Hi": 13,
    "Recall": 10, "Precision": 11, "F1": 8, "F5": 8,
    "Applied UF": 11, "% True Above AET": 16, "% Reported Above AET": 20,
}

RATE_COLS = {
    "FN Rate", "FN 95% PI Lo", "FN 95% PI Hi",
    "FP Rate", "FP 95% PI Lo", "FP 95% PI Hi",
    "Recall", "Precision", "F1", "F5",
    "% True Above AET", "% Reported Above AET", "AET (\u00b5g/mL)",
}


def write_sheet(ws, df_sheet):
    df_out = df_sheet.reindex(columns=COLS)
    for c_idx, col_name in enumerate(COLS, start=1):
        cell = ws.cell(row=1, column=c_idx, value=col_name)
        cell.font      = HDR_FONT
        cell.fill      = PI_FILL if col_name in PI_COLS else HDR_FILL
        cell.alignment = HDR_ALIGN
        cell.border    = CELL_BDR
    for r_idx, (_, row) in enumerate(df_out.iterrows(), start=2):
        fill = ALT_FILL if r_idx % 2 == 0 else None
        for c_idx, col_name in enumerate(COLS, start=1):
            val  = row[col_name]
            cell = ws.cell(row=r_idx, column=c_idx,
                           value=None if (isinstance(val, float) and pd.isna(val)) else val)
            cell.border    = CELL_BDR
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if fill:
                cell.fill = fill
            if col_name in ("FN SD", "FP SD"):
                cell.number_format = NUM_FMT_SD
            elif col_name in RATE_COLS:
                cell.number_format = NUM_FMT
    for c_idx, col_name in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = COL_WIDTHS.get(col_name, 14)
    ws.freeze_panes    = "A2"
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
wb.remove(wb.active)

for sheet_name, df_sheet in SHEET_MAP.items():
    ws = wb.create_sheet(sheet_name)
    write_sheet(ws, df_sheet)
    print(f"  Sheet '{sheet_name}': {len(df_sheet)} rows")

wb.save(XLSX_PATH)
print(f"\nDone - Results_by_scenario.xlsx written to:\n  {XLSX_PATH}")
")
